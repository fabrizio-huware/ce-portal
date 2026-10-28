"""Costi vari negli export: Excel con formule vive, CSV, PDF, vista ridotta e divisione per business unit."""

import re
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from pypdf import PdfReader

from app.exports.pdf_export import ce_full_pdf
from app.exports.xlsx_export import ce_full_xlsx
from app.schemas.ce import CEDetail
from app.services import ce as svc
from app.services.unit_division import (
    EXTRAS,
    business_unit_division,
    business_unit_revenue,
)
from tests.ce_helpers import approved_ce, create, save, simple_content
from tests.xlsx_helpers import formula_errors, needs_libreoffice, recalculated

URL = "/api/v1/ce"
NOW = datetime(2026, 10, 6, 12, 30, tzinfo=UTC)
SECRET = "RISERVATO-NON-USCIRE"
EXTRAS_IN = [
    {"description": f"Viaggi {SECRET}", "cost": "1000.00", "markup_pct": "20.00"},
    {"description": "Licenze", "cost": "333.33", "markup_pct": "15.50"},
]  # ricavi 1.200,00 + 385,00 = 1.585,00; costi 1.333,33


def content(env, revision, extras=EXTRAS_IN):
    body = simple_content(env, revision)  # 16 ore di Senior: ricavi 1.700, costi 660
    body["extra_costs"] = extras
    return body


def text_of(pdf: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)


def detail_of(api, env, ce_id) -> CEDetail:
    return CEDetail.model_validate(api.get(f"{URL}/{ce_id}", headers=env.h(env.admin)).json())


@pytest.fixture
def ce_id(api, env):
    created = create(api, env, code="PS-EXTRA-EXP")
    resp = save(api, env, created["ce"]["id"], content(env, 1))
    assert resp.status_code == 200, resp.text
    return created["ce"]["id"]


def export(api, env, ce_id, fmt, user=None, variant="full"):
    resp = api.get(
        f"{URL}/{ce_id}/export?format={fmt}&variant={variant}", headers=env.h(user or env.presale)
    )
    assert resp.status_code == 200, resp.text
    return resp.content


# ------------------------------------------------------------------ Excel
@needs_libreoffice
def test_the_excel_formulas_recalculate_to_the_engine_numbers(api, env, ce_id):
    wb = recalculated(ce_full_xlsx(detail_of(api, env, ce_id), NOW))
    assert formula_errors(wb) == []
    ex = wb["Costi vari"]
    assert ex["D2"].value == pytest.approx(1200.00) and ex["D3"].value == pytest.approx(385.00)
    assert ex["B4"].value == pytest.approx(1333.33) and ex["D4"].value == pytest.approx(1585.00)
    s = wb["Riepilogo"]
    assert s["B22"].value == pytest.approx(1585.00) and s["C22"].value == pytest.approx(1333.33)
    assert s["B23"].value == pytest.approx(3285.00) and s["C23"].value == pytest.approx(1993.33)


@needs_libreoffice
def test_changing_an_extra_row_in_excel_changes_the_totals(api, env, ce_id):
    def edit(wb):
        wb["Costi vari"]["B2"] = 2000

    wb = recalculated(ce_full_xlsx(detail_of(api, env, ce_id), NOW), edit)
    assert wb["Costi vari"]["D2"].value == pytest.approx(2400.00)
    assert wb["Riepilogo"]["B23"].value == pytest.approx(4485.00)


def test_a_ce_without_extras_has_no_extra_sheet(api, env):
    created = create(api, env, code="PS-EXTRA-NONE")
    save(api, env, created["ce"]["id"], simple_content(env, 1))
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW)))
    assert "Costi vari" not in wb.sheetnames


# ------------------------------------------------------------------ CSV
def test_the_csv_has_one_row_per_extra_with_the_markup(api, env, ce_id):
    raw = export(api, env, ce_id, "csv").decode("utf-8-sig")
    lines = raw.strip().splitlines()
    assert lines[0].endswith("Ricarico (%)")
    extra = [ln for ln in lines if ";Costi vari;" in ln or ",Costi vari," in ln]
    assert len(extra) == 2
    assert "1200,00" in extra[0] and "1000,00" in extra[0] and "200,00" in extra[0]
    assert "20,00" in extra[0].split(";")[-1] or "20" in extra[0].split(";")[-1]


# ------------------------------------------------------------------ PDF completo
def test_the_full_pdf_shows_the_extra_costs(api, env, ce_id):
    text = text_of(ce_full_pdf(detail_of(api, env, ce_id), NOW))
    for expected in ("Viaggi " + SECRET, "Licenze", "Ricarico (%)", "Totale riga (€)", "Totale costi vari",
                     "3.285,00"):  # fmt: skip
        assert expected in text
    assert re.search(r"Costi vari\s+1\.585,00\s+1\.333,33", text)  # la riga in "Ricavi e costi"
    assert text.count("Costi vari") >= 2  # e la sezione con le righe


# ------------------------------------------------------------------ vista ridotta
def test_the_reduced_views_show_one_costi_vari_row_and_never_the_details(api, env, ce_id):
    approved = approved_ce(api, env, "PS-EXTRA-APP", content=content(env, 1))
    new_id = approved["ce"]["id"]
    for fmt in ("csv", "pdf", "xlsx"):
        viewer = api.get(f"{URL}/summaries/{new_id}/export?format={fmt}", headers=env.h(env.viewer))
        assert viewer.status_code == 200
        raw = viewer.content
        blob = raw.decode("utf-8-sig") if fmt == "csv" else (text_of(raw) if fmt == "pdf" else "")
        if fmt == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(BytesIO(raw))
            blob = " ".join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value)
        assert "Costi vari" in blob and any(v in blob for v in ("1.585,00", "1585,00", "1585"))
        for secret in (
            SECRET,
            "Licenze",
            "Viaggi",
            "1.333,33",
            "1333,33",
            "333,33",
            "15,5",
            "20,00%",
        ):
            assert secret not in blob, (fmt, secret)


def test_the_editor_summary_pdf_has_the_extras_without_the_details(api, env, ce_id):
    text = text_of(export(api, env, ce_id, "pdf", variant="summary"))
    assert "Costi vari" in text and "1.585,00" in text
    assert SECRET not in text and "Licenze" not in text and "1.333,33" not in text


def test_the_editor_summary_csv_has_the_extras_without_the_details(api, env, ce_id):
    raw = export(api, env, ce_id, "csv", variant="summary").decode("utf-8-sig")
    assert "Costi vari" in raw and "1585,00" in raw
    assert SECRET not in raw and "Licenze" not in raw and "1333,33" not in raw


def test_the_viewer_json_has_the_extras_as_a_phase_with_only_the_revenue(api, env):
    approved = approved_ce(api, env, "PS-EXTRA-VJ", content=content(env, 1))
    resp = api.get(f"{URL}/summaries/{approved['ce']['id']}", headers=env.h(env.viewer))
    assert resp.status_code == 200, resp.text
    phases = {p["name"]: p["revenue"] for p in resp.json()["phases"]}
    assert Decimal(phases["Costi vari"]) == Decimal("1585.00")
    assert SECRET not in resp.text and "Licenze" not in resp.text and "markup" not in resp.text
    assert svc  # il modulo resta importato per i test che lo usano


# ------------------------------------------------------------------ divisione
def test_the_division_has_an_extras_row_and_adds_up_to_the_totals(api, env, session, ce_id):
    detail = detail_of(api, env, ce_id)
    rows = business_unit_division(session, detail)
    assert rows[-1].label == EXTRAS and rows[-1].kind == "extras"
    assert rows[-1].days == 0 and rows[-1].revenue == Decimal("1585.00")
    assert rows[-1].cost == Decimal("1333.33")
    total = detail.calculation.total
    assert sum(r.revenue for r in rows) == total.revenue
    assert sum(r.cost for r in rows) == total.cost
    assert sum(r.days for r in rows) == total.days


def test_the_revenue_by_unit_has_the_extras_last_and_adds_up(api, env, session, ce_id):
    detail = detail_of(api, env, ce_id)
    rows = business_unit_revenue(session, detail)
    assert rows[-1].label == EXTRAS and rows[-1].revenue == Decimal("1585.00")
    assert sum(r.revenue for r in rows) == detail.calculation.total.revenue
    assert sum(r.share for r in rows) == pytest.approx(1)


def test_a_ce_with_only_extras_still_divides(api, env, session):
    created = create(api, env, code="PS-EXTRA-ONLY")
    body = simple_content(env, 1)
    body["phases"] = []
    body["extra_costs"] = EXTRAS_IN
    resp = save(api, env, created["ce"]["id"], body)
    assert resp.status_code == 200, resp.text
    detail = detail_of(api, env, created["ce"]["id"])
    assert [r.label for r in business_unit_division(session, detail)] == [EXTRAS]
    assert [r.label for r in business_unit_revenue(session, detail)] == [EXTRAS]


@needs_libreoffice
def test_the_excel_recalculates_a_negative_markup(api, env):
    created = create(api, env, code="PS-EXTRA-NEG")
    rows = [{"description": "Sconto", "cost": "1000.00", "markup_pct": "-25.00"}]
    assert save(api, env, created["ce"]["id"], content(env, 1, rows)).status_code == 200
    wb = recalculated(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW))
    assert formula_errors(wb) == []
    ex = wb["Costi vari"]
    assert ex["D2"].value == pytest.approx(750.00) and ex["E2"].value == pytest.approx(-250.00)
    assert wb["Riepilogo"]["B22"].value == pytest.approx(750.00)
    assert wb["Riepilogo"]["B23"].value == pytest.approx(2450.00)
