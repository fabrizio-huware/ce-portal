"""Portfolio e Carico risorse con l'esito: Close Won, approvati senza esito, Close Lost e pipeline."""

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from tests.dash_helpers import world

D = Decimal
PORTFOLIO = "/api/v1/dashboard/portfolio"
RESOURCES = "/api/v1/dashboard/resources?date_from=2026-01-01&date_to=2026-12-31"


def _outcome(api, env, ce, value):
    resp = api.put(f"/api/v1/ce/{ce}/outcome", headers=env.h(env.admin), json={"outcome": value})
    assert resp.status_code == 200, resp.text


def _portfolio(api, env, query=""):
    resp = api.get(f"{PORTFOLIO}?{query}", headers=env.h(env.presale))
    assert resp.status_code == 200, resp.text
    return resp.json()


def _outcomes(api, env, w):
    """A (23.000) Close Won, B (4.250) Close Lost; D (3.625) resta approvato senza esito; C e F in pipeline."""
    _outcome(api, env, w["ids"]["A"], "won")
    _outcome(api, env, w["ids"]["B"], "lost")


def _revenue(totals):
    return D(totals["revenue"])


# ------------------------------------------------------------------ Portfolio
def test_the_approved_series_is_split_by_outcome(api, env, session):
    w = world(api, env, session)
    before = _portfolio(api, env)  # nessun esito: come prima
    assert (before["won"]["count"], before["lost"]["count"]) == (0, 0)
    assert before["confirmed"] == before["approved"]
    _outcomes(api, env, w)
    after = _portfolio(api, env)
    assert (after["approved"]["count"], after["won"]["count"], after["lost"]["count"]) == (1, 1, 1)
    assert (_revenue(after["approved"]), _revenue(after["won"]), _revenue(after["lost"])) == (
        D("3625"),
        D("23000"),
        D("4250"),
    )
    assert (
        _revenue(after["confirmed"]) == D("26625") and after["confirmed"]["count"] == 2
    )  # approvati senza esito + Close Won
    assert after["pipeline"] == before["pipeline"]  # la pipeline non cambia
    assert _revenue(after["approved"]) + _revenue(after["won"]) + _revenue(
        after["lost"]
    ) == _revenue(before["approved"])


def test_a_close_lost_is_not_counted_as_approved_anywhere_but_is_listed(api, env, session):
    _outcomes(api, env, world(api, env, session))
    body = _portfolio(api, env)
    items = {i["code"]: i for i in body["items"]}
    assert (
        items["PS-DASH-A"]["scope"],
        items["PS-DASH-B"]["scope"],
        items["PS-DASH-D"]["scope"],
    ) == ("won", "lost", "approved")
    assert items["PS-DASH-C"]["scope"] == "pipeline"
    rows = {(s["label"], s["scope"]): s for s in body["by_status"]}
    assert (
        ("Close Won", "won") in rows
        and ("Close Lost", "lost") in rows
        and ("Approvato", "approved") in rows
    )
    assert D(rows[("Close Lost", "lost")]["revenue"]) == D("4250")
    beta = next(g for g in body["by_client"] if g["label"] == "Beta Srl")
    assert (beta["lost"]["count"], _revenue(beta["lost"]), beta["approved"]["count"]) == (
        1,
        D("4250"),
        1,
    )
    alfa = next(g for g in body["by_client"] if g["label"] != "Beta Srl")
    assert _revenue(alfa["won"]) == D("23000") and alfa["lost"]["count"] == 0


def test_the_monthly_view_has_close_won_and_never_close_lost(api, env, session):
    w = world(api, env, session)
    before = _portfolio(api, env)
    _outcomes(api, env, w)
    after = _portfolio(api, env)

    def total(body):
        return sum(
            D(m[k]["revenue"]) for m in body["by_month"] for k in ("approved", "pipeline", "won")
        )

    won_months = sum(D(m["won"]["revenue"]) for m in after["by_month"])
    assert won_months > 0  # le serie mensili hanno i Close Won
    assert total(after) < total(before)  # e non hanno i Close Lost: il totale scende
    lost_in_months = total(before) - total(after)
    assert abs(lost_in_months - D("4250")) < D(
        "0.05"
    )  # esattamente i ricavi mensili di B (arrotondamenti a parte)


def test_the_filters_still_work_with_outcomes(api, env, session):
    _outcomes(api, env, world(api, env, session))
    only_alfa = _portfolio(api, env, f"client_id={env.client.id}")
    assert only_alfa["won"]["count"] == 1 and only_alfa["lost"]["count"] == 0
    no_pipeline = _portfolio(api, env, "include_pipeline=false")
    assert no_pipeline["pipeline"]["count"] == 0 and no_pipeline["lost"]["count"] == 1


def test_the_portfolio_exports_have_the_new_rows_and_columns(api, env, session):
    _outcomes(api, env, world(api, env, session))
    csv = api.get(f"{PORTFOLIO}/export?format=csv", headers=env.h(env.presale)).text
    assert "Close Won" in csv and "Close Lost" in csv
    xlsx = api.get(f"{PORTFOLIO}/export?format=xlsx", headers=env.h(env.presale))
    wb = load_workbook(BytesIO(xlsx.content))
    summary = {row[0].value: row for row in wb["Riepilogo"].iter_rows(min_row=5, max_row=8)}
    assert set(summary) == {"Approvati", "Pipeline", "Close Won", "Close Lost"}
    assert summary["Close Won"][2].value == D("23000") and summary["Close Lost"][2].value == D(
        "4250"
    )
    assert (
        wb["Per cliente"][1][-1].value == "Ricavi Close Lost (€)"
    )  # le colonne nuove sono in coda
    assert wb["Per mese"][1][-1].value == "Giornate Close Won"


# ------------------------------------------------------------------ Carico risorse
def _ada_days(api, env):
    body = api.get(RESOURCES, headers=env.h(env.admin)).json()
    ada = next(e for e in body["employees"] if e["name"].startswith("Ada"))
    return D(ada["total_days"]), body["ces_count"]


def test_a_close_lost_leaves_the_resources_and_a_close_won_stays(api, env, session):
    w = world(api, env, session)
    base_days, base_count = _ada_days(api, env)  # A (10 g) e B (5 g)
    assert base_days == D("15")
    _outcome(api, env, w["ids"]["A"], "won")
    assert _ada_days(api, env) == (base_days, base_count)  # un Close Won pesa come un approvato
    _outcome(api, env, w["ids"]["B"], "lost")
    days, count = _ada_days(api, env)
    assert days == D("10") and count == base_count - 1  # B non pesa più sulle risorse
    _outcome(api, env, w["ids"]["B"], None)
    assert _ada_days(api, env) == (base_days, base_count)  # e torna quando si toglie l'esito


def test_a_close_won_counts_as_approved_not_as_pipeline_in_the_resources(api, env, session):
    w = world(api, env, session)
    _outcome(api, env, w["ids"]["A"], "won")
    body = api.get(f"{RESOURCES}&include_pipeline=true", headers=env.h(env.admin)).json()
    ada = next(e for e in body["employees"] if e["name"].startswith("Ada"))
    approved = sum(D(c["days_approved"]) for c in ada["cells"])
    pipeline = sum(D(c["days_pipeline"]) for c in ada["cells"])
    assert approved == D("15") and pipeline == D("0")
