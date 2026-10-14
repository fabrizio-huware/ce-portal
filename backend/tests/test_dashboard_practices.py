"""Carico risorse per practice.

Percentuale di una practice = giorni della practice / (persone che hanno impegno nel periodo x giorni lavorativi del
mese). La practice di un collaboratore è quella attuale, per tutto il periodo. Settembre 2026: 22 giorni lavorativi.
"""

from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from tests.ce_helpers import add_closures
from tests.dash_helpers import build_ce, make_employee
from tests.test_export_api import csv_rows

URL = "/api/v1/dashboard/resources"
D = Decimal
SEP = "date_from=2026-09-01&date_to=2026-09-30"


def get(api, env, query=SEP):
    resp = api.get(f"{URL}?{query}", headers=env.h(env.presale))
    assert resp.status_code == 200, resp.text
    return resp.json()


def row(body, kind, name):
    return next(r for r in body[kind] if r["name"] == name)


def load(
    api,
    env,
    code,
    who,
    hours,
    *,
    start="2026-09-01",
    end="2026-09-30",
    state="approved",
    profile="Senior",
):
    build_ce(api, env, code, start=start, end=end, lines=[(profile, hours, who)], state=state)


@pytest.fixture
def people(session, env):
    """Cloud: Ada e Bruno. Data: Carla. «Da assegnare»: Dario."""
    return {
        "ada": make_employee(session, env, "Ada", "L", practice="Cloud"),
        "bruno": make_employee(session, env, "Bruno", "N", practice="Cloud"),
        "carla": make_employee(session, env, "Carla", "M", practice="Data"),
        "dario": make_employee(session, env, "Dario", "P"),
    }


def test_practice_rows_sum_the_days_of_their_people_over_persons_times_capacity(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)  # 11 giorni = 50%
    load(api, env, "PS-PR-2", people["bruno"], 176)  # 22 giorni = 100%
    load(api, env, "PS-PR-3", people["carla"], 264)  # 33 giorni = 150%
    body = get(api, env)
    cloud, data = row(body, "practices", "Cloud"), row(body, "practices", "Data")
    c = cloud["cells"][0]
    assert cloud["kind"] == "practice" and cloud["persons"] == 2
    assert (D(c["days"]), D(c["hours"]), D(c["fte"]), c["overloaded"]) == (
        D(33),
        D(264),
        D("0.75"),
        False,
    )
    assert D(cloud["total_days"]) == 33 and cloud["overloaded_months"] == 0
    d = data["cells"][0]
    assert data["persons"] == 1 and (D(d["days"]), D(d["fte"]), d["overloaded"]) == (
        D(33),
        D("1.5"),
        True,
    )
    assert data["overloaded_months"] == 1
    assert [r["name"] for r in body["practices"]] == [
        "Cloud",
        "Data",
    ]  # senza i colleghi senza impegno


def test_someone_with_load_in_the_period_counts_in_every_month_of_it(api, env, people):
    load(api, env, "PS-PR-A", people["ada"], 176)  # solo settembre
    load(
        api, env, "PS-PR-B", people["bruno"], 176, start="2026-10-01", end="2026-10-31"
    )  # solo ottobre
    body = get(api, env, "date_from=2026-09-01&date_to=2026-10-31")
    cloud = row(body, "practices", "Cloud")
    assert cloud["persons"] == 2
    caps = [m["capacity_days"] for m in body["months"]]
    for cell, cap in zip(cloud["cells"], caps, strict=True):
        assert D(cell["fte"]) == D(cell["days"]) / (
            2 * cap
        )  # entrambi contano a settembre e a ottobre
    assert D(cloud["cells"][0]["fte"]) == D("0.5")  # settembre: 22 giorni su 2 persone x 22


def test_people_without_load_in_the_period_are_not_counted(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)  # Bruno, collega di Ada, non ha impegno
    body = get(api, env)
    cloud = row(body, "practices", "Cloud")
    assert cloud["persons"] == 1 and D(cloud["cells"][0]["fte"]) == D("0.5")
    assert [r["name"] for r in body["employees"]] == ["Ada L"]


def test_a_practice_with_nobody_working_is_not_listed(api, env, people):
    load(api, env, "PS-PR-1", people["carla"], 88)
    assert [r["name"] for r in get(api, env)["practices"]] == [
        "Data"
    ]  # Cloud e «Da assegnare» non compaiono


def test_pipeline_is_excluded_by_default_and_counted_with_the_flag(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)
    load(api, env, "PS-PR-2", people["bruno"], 176, state="draft")
    default = row(get(api, env), "practices", "Cloud")
    assert default["persons"] == 1 and D(default["cells"][0]["days"]) == 11
    full = row(get(api, env, SEP + "&include_pipeline=true"), "practices", "Cloud")
    c = full["cells"][0]
    assert full["persons"] == 2
    assert (D(c["days_approved"]), D(c["days_pipeline"]), D(c["days"]), D(c["fte"])) == (
        D(11),
        D(22),
        D(33),
        D("0.75"),
    )


def test_a_person_on_several_ces_counts_once(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 44)
    load(api, env, "PS-PR-2", people["ada"], 44)
    cloud = row(get(api, env), "practices", "Cloud")
    assert cloud["persons"] == 1 and D(cloud["cells"][0]["days"]) == 11


def test_employee_rows_carry_their_practice(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)
    load(api, env, "PS-PR-2", people["dario"], 88)
    body = get(api, env)
    assert row(body, "employees", "Ada L")["practice_name"] == "Cloud"
    assert row(body, "employees", "Dario P")["practice_name"] == "Da assegnare"
    assert row(body, "practices", "Da assegnare")["persons"] == 1


def test_the_practice_is_the_current_one_for_the_whole_period(api, env, session, people):
    load(api, env, "PS-PR-1", people["ada"], 88)
    before = get(api, env)
    assert [r["name"] for r in before["practices"]] == ["Cloud"]
    data = people["carla"].practice_id
    people[
        "ada"
    ].practice_id = data  # Ada passa a Data: vale per tutti i mesi, anche quelli già passati
    session.flush()
    after = get(api, env)
    assert [r["name"] for r in after["practices"]] == ["Data"] and row(after, "employees", "Ada L")[
        "practice_name"
    ] == "Data"


def test_exactly_100_percent_on_average_is_not_overload_and_just_above_is(api, env, session):
    cloud = [make_employee(session, env, n, "C", practice="Cloud") for n in ("Ada", "Bruno")]
    data = [make_employee(session, env, n, "D", practice="Data") for n in ("Carla", "Dario")]
    for i, (who, hours) in enumerate(zip(cloud + data, (176, 176, 176, 184), strict=True)):
        load(api, env, f"PS-PR-{i}", who, hours)
    body = get(api, env)
    exact, above = (
        row(body, "practices", "Cloud")["cells"][0],
        row(body, "practices", "Data")["cells"][0],
    )
    assert D(exact["fte"]) == 1 and exact["overloaded"] is False
    assert D(above["fte"]) > 1 and above["overloaded"] is True  # (22 + 23) / 44


def test_closures_lower_the_capacity_of_the_practice_too(api, env, session, people):
    from datetime import date

    add_closures(session, date(2026, 9, 14), date(2026, 9, 15))  # due giorni feriali: capacità 20
    load(api, env, "PS-PR-1", people["ada"], 88)  # 11 giorni
    body = get(api, env)
    assert body["months"][0]["capacity_days"] == 20
    assert D(row(body, "practices", "Cloud")["cells"][0]["fte"]) == D("0.55")


def test_filtering_by_practice_keeps_only_its_people_and_lines(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)
    load(api, env, "PS-PR-2", people["carla"], 264)
    build_ce(
        api,
        env,
        "PS-PR-3",
        start="2026-09-01",
        end="2026-09-30",
        lines=[("Senior", 40, None)],
        state="approved",
    )
    everything = get(api, env)
    assert D(row(everything, "profiles", "Senior")["unassigned_days"]) == 5
    body = get(api, env, SEP + f"&practice_id={people['ada'].practice_id}")
    assert [r["name"] for r in body["employees"]] == ["Ada L"] and [
        r["name"] for r in body["practices"]
    ] == ["Cloud"]
    senior = row(body, "profiles", "Senior")
    assert (
        D(senior["cells"][0]["days"]) == 11 and D(senior["unassigned_days"]) == 0
    )  # le righe senza persona sono fuori


def test_an_unknown_practice_filter_is_harmless_and_no_data_means_no_rows(api, env, people):
    body = get(api, env, SEP + f"&practice_id={uuid4()}")
    assert (body["employees"], body["practices"], body["profiles"]) == ([], [], [])
    assert get(api, env)["practices"] == []


# ---------------------------------------------------------------------- esportazioni
def _export(api, env, **params):
    return api.get("/api/v1/dashboard/resources/export", headers=env.h(env.presale),
                   params={"date_from": "2026-09-01", "date_to": "2026-09-30", **params})  # fmt: skip


@pytest.fixture
def loaded(api, env, people):
    load(api, env, "PS-PR-1", people["ada"], 88)
    load(api, env, "PS-PR-2", people["bruno"], 176)
    load(api, env, "PS-PR-3", people["carla"], 264)
    return people


def test_csv_keeps_the_old_columns_and_adds_practice_and_persons_at_the_end(api, env, loaded):
    rows = csv_rows(_export(api, env, format="csv"))
    assert rows[0][:11] == ["Tipo", "Nome", "Profilo", "Mese", "Giorni approvati", "Giorni pipeline", "Giorni totali",
                            "Ore", "Capacità (giorni)", "FTE (%)", "Sovraccarico"]  # fmt: skip
    assert rows[0][11:] == ["Practice", "Persone"]
    col = {n: i for i, n in enumerate(rows[0])}
    cloud = next(r for r in rows[1:] if r[col["Tipo"]] == "Practice" and r[col["Nome"]] == "Cloud")
    assert (cloud[col["Practice"]], cloud[col["Persone"]], cloud[col["Giorni totali"]]) == (
        "Cloud",
        "2",
        "33,00",
    )
    assert (
        cloud[col["Capacità (giorni)"]] == "44" and cloud[col["FTE (%)"]] == "75,00"
    )  # 22 giorni x 2 persone
    ada = next(r for r in rows[1:] if r[col["Nome"]] == "Ada L")
    assert (
        ada[col["Tipo"]] == "Collaboratore"
        and ada[col["Practice"]] == "Cloud"
        and ada[col["Persone"]] == ""
    )


def test_excel_has_a_practice_sheet_and_the_practice_columns_in_the_data_sheet(api, env, loaded):
    wb = load_workbook(BytesIO(_export(api, env).content))
    assert wb.sheetnames == ["Collaboratori", "Practice", "Profili", "Dati"]
    ws = wb["Practice"]
    assert (ws["A1"].value, ws["B1"].value) == ("Practice", "Persone con impegno")
    assert (ws["A2"].value, ws["B2"].value) == ("Cloud", 2) and ws["C2"].value == pytest.approx(
        0.75
    )
    assert (ws["A3"].value, ws["B3"].value) == ("Data", 1) and ws["C3"].value == pytest.approx(1.5)
    assert ws["C3"].fill.fgColor.rgb.endswith("FDE2E1") and not ws["C2"].fill.fgColor.rgb.endswith(
        "FDE2E1"
    )
    assert any(
        "persone che hanno impegno" in str(c.value) for r in ws.iter_rows() for c in r if c.value
    )
    data = wb["Dati"]
    header = [c.value for c in data[1]]
    assert header[-2:] == ["Practice", "Persone"]
    practice_rows = [r for r in data.iter_rows(min_row=2, values_only=True) if r[0] == "Practice"]
    assert {(r[1], r[11], r[12]) for r in practice_rows} == {
        ("Cloud", "Cloud", 2),
        ("Data", "Data", 1),
    }


def test_the_export_can_be_filtered_by_practice(api, env, loaded):
    rows = csv_rows(_export(api, env, format="csv", practice_id=str(loaded["carla"].practice_id)))
    col = {n: i for i, n in enumerate(rows[0])}
    assert {r[col["Nome"]] for r in rows[1:] if r[col["Tipo"]] != "Profilo"} == {"Carla M", "Data"}
