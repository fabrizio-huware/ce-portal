"""Sicurezza e formati degli export: formule, nomi di file, numeri e date italiani."""

from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.exports.common import (
    csv_num,
    fmt_date,
    fmt_datetime,
    fmt_money,
    fmt_num,
    fmt_pct,
    neutralize,
    safe_filename,
)
from app.exports.xlsx_export import formula, put


@pytest.mark.parametrize(
    "hostile",
    ["=1+1", "+39 333", "-5% sconto", "@SUM(A1)", "\tTab", "\rA-capo", '=HYPERLINK("http://x")'],
)
def test_text_that_could_run_as_a_formula_is_neutralised_in_csv(hostile):
    assert neutralize(hostile) == "'" + hostile


@pytest.mark.parametrize("fine", ["Analisi", "10 giorni", "Progetto -X", "", "a=b"])
def test_normal_text_is_left_untouched(fine):
    assert neutralize(fine) == fine


def test_numbers_and_non_strings_are_never_modified():
    assert (
        neutralize(-5) == -5
        and neutralize(None) is None
        and neutralize(Decimal("-1.5")) == Decimal("-1.5")
    )


def test_excel_text_cells_stay_text_even_if_they_start_like_a_formula():
    ws = Workbook().active
    for ref, text in (("A1", "=1+1"), ("A2", "+39 333"), ("A3", "-5"), ("A4", "@x")):
        put(ws, ref, text)
        assert ws[ref].data_type == "s" and ws[ref].value == text and ws[ref].quotePrefix is True
    put(ws, "A5", "Normale")
    assert ws["A5"].data_type == "s" and not ws["A5"].quotePrefix
    formula(ws, "B1", "SUM(1,2)")
    assert ws["B1"].data_type == "f"  # le formule vere restano formule


@pytest.mark.parametrize(
    "parts,ext,expected",
    [
        (("PS-BONGA-AI-PJT", "v1"), "xlsx", "PS-BONGA-AI-PJT_v1.xlsx"),
        (("PS-BONGA", "v2", "riepilogo"), "pdf", "PS-BONGA_v2_riepilogo.pdf"),
        (("../../etc/passwd",), "csv", "etc_passwd.csv"),
        (("CE con spazi/e\\slash",), "csv", "CE_con_spazi_e_slash.csv"),
        (("Società è così",), "xlsx", "Societa_e_cosi.xlsx"),
        (('a"b;c\r\nSet-Cookie: x',), "csv", "a_b_c_Set-Cookie_x.csv"),
        (("",), "csv", "export.csv"),
        (("😀😀",), "pdf", "export.pdf"),
    ],
)
def test_file_names_are_safe_ascii(parts, ext, expected):
    assert safe_filename(*parts, ext=ext) == expected


def test_file_names_are_bounded_in_length():
    assert len(safe_filename("x" * 500, ext="pdf")) <= 124


def test_italian_number_formats():
    assert csv_num(Decimal("1234.5")) == "1234,50" and csv_num(Decimal("0.125"), 4) == "0,1250"
    assert csv_num(None) == "" and csv_num(Decimal("-3.456")) == "-3,46"
    assert fmt_num(Decimal("1234567.5")) == "1.234.567,50" and fmt_num(None) == "n/d"
    assert fmt_money(Decimal("47162.5")) == "47.162,50 €" and fmt_money(None) == "n/d"
    assert fmt_pct(Decimal("0.571694")) == "57,17%" and fmt_pct(None) == "n/d"
    assert fmt_num(Decimal("0.005")) == "0,01"  # mezzo centesimo per eccesso, come nel motore


def test_dates_and_times_are_shown_in_italian_local_time():
    assert fmt_date(date(2026, 1, 5)) == "05/01/2026" and fmt_date(None) == ""
    assert (
        fmt_datetime(datetime(2026, 10, 6, 12, 30, tzinfo=UTC)) == "06/10/2026 14:30"
    )  # ora legale
    assert (
        fmt_datetime(datetime(2026, 1, 6, 12, 30, tzinfo=UTC)) == "06/01/2026 13:30"
    )  # ora solare
    assert (
        fmt_datetime(datetime(2026, 12, 31, 23, 30, tzinfo=UTC)) == "01/01/2027 00:30"
    )  # cambia anche il giorno
    assert fmt_datetime(None) == ""
