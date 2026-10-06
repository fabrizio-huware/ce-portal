"""Excel completo: le formule ricalcolate da LibreOffice devono dare i numeri del motore."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.exports.xlsx_export import ce_full_xlsx
from app.schemas.ce import CEDetail
from tests.ce_helpers import create, percent_ce_content, real_ce_content, save, simple_content
from tests.engine_helpers import FIXTURE
from tests.xlsx_helpers import formula_errors, needs_libreoffice, recalculated

URL = "/api/v1/ce"
NOW = datetime(2026, 10, 6, 12, 30, tzinfo=UTC)


def detail_of(api, env, ce_id) -> CEDetail:
    return CEDetail.model_validate(api.get(f"{URL}/{ce_id}", headers=env.h(env.presale)).json())


@pytest.fixture
def real(api, env):
    """Il CE reale (468 ore): ricavi 47.162,50, costi 20.200,00, 58,5 giornate."""
    created = create(api, env, code="PS-XLS-REAL")
    save(api, env, created["ce"]["id"], real_ce_content(env, 1))
    return detail_of(api, env, created["ce"]["id"])


@needs_libreoffice
def test_summary_formulas_recalculate_to_the_engine_numbers(real):
    wb = recalculated(ce_full_xlsx(real, NOW))
    assert formula_errors(wb) == []
    s = wb["Riepilogo"]
    assert s["B19"].value == pytest.approx(47162.50) and s["C19"].value == pytest.approx(
        20200.00
    )  # interni
    assert s["B20"].value == 0 and s["B21"].value == 0  # esterni e contingency
    assert s["B22"].value == pytest.approx(47162.50) and s["C22"].value == pytest.approx(20200.00)
    assert s["D22"].value == pytest.approx(26962.50)
    assert s["E22"].value == pytest.approx(0.571694, abs=1e-6) and s["F22"].value == pytest.approx(
        0.428306, abs=1e-6
    )
    assert s["G22"].value == pytest.approx(58.5)
    assert s["B25"].value == pytest.approx(47162.50)  # prezzo progetto
    assert s["B27"].value == pytest.approx(47162.50)  # prezzo minimo (nessuno sconto)
    assert s["B29"].value == pytest.approx(10) and s["B30"].value == pytest.approx(48.5)
    assert s["B31"].value == pytest.approx(11.7)
    assert s["B34"].value == pytest.approx(806.1966, abs=1e-3)
    assert s["B36"].value == "n/d" and s["B37"].value == "n/d"  # nessun prezzo firmato


@needs_libreoffice
def test_profile_breakdown_matches_the_sheet(real):
    wb = recalculated(ce_full_xlsx(real, NOW))
    rows = {r[0].value: r for r in wb["Profili"].iter_rows(min_row=2, max_row=10)}
    assert rows["Specialist"][2].value == pytest.approx(292) and rows["Specialist"][
        5
    ].value == pytest.approx(26462.50)
    assert rows["Senior"][6].value == pytest.approx(4620.00) and rows["Manager"][
        5
    ].value == pytest.approx(4800.00)
    assert rows["Esterni"][2].value == 0
    assert wb["Profili"]["C11"].value == pytest.approx(468)


@needs_libreoffice
def test_phase_sheet_and_totals(real):
    wb = recalculated(ce_full_xlsx(real, NOW))
    ph = wb["Fasi"]
    assert [ph[f"A{r}"].value for r in range(2, 7)] == ["PM", "AN", "SD", "TA", "TST"]
    assert ph["B2"].value == pytest.approx(80) and ph["D2"].value == pytest.approx(
        10375.0
    )  # PM: 10 giornate
    assert ph["D4"].value == pytest.approx(2550.0)  # SD
    assert ph["D5"].value == pytest.approx(20200.0)  # TA
    assert ph["D11"].value == pytest.approx(47162.5) and ph["B11"].value == pytest.approx(
        468
    )  # totale


@needs_libreoffice
def test_monthly_staffing_matches_the_engine(real):
    wb = recalculated(ce_full_xlsx(real, NOW))
    st = wb["Staffing mensile"]
    cal = real.calculation
    n = len(cal.months)
    assert [st.cell(row=4, column=6 + i).value for i in range(n)] == [
        17,
        20,
        22,
        22,
    ]  # giorni lavorativi
    for i, month in enumerate(cal.monthly):
        col = 6 + i
        assert st.cell(row=4, column=col).value == month.work_days
    first_total_row = 6 + len(cal.lines) + 2
    for i, month in enumerate(cal.monthly):
        col = 6 + i
        assert st.cell(row=first_total_row, column=col).value == pytest.approx(
            float(month.days), abs=1e-6
        )
        assert st.cell(row=first_total_row + 2, column=col).value == pytest.approx(
            float(month.fte), abs=1e-5
        )
        assert st.cell(row=first_total_row + 3, column=col).value == pytest.approx(
            float(month.revenue), abs=0.02
        )
        assert st.cell(row=first_total_row + 4, column=col).value == pytest.approx(
            float(month.cost), abs=0.02
        )
    total_col = 6 + n
    assert st.cell(row=first_total_row, column=total_col).value == pytest.approx(58.5)


@needs_libreoffice
def test_formulas_are_live_changing_inputs_changes_the_results(real):
    """Cambiando ore, sconto, contingency, prezzo firmato e tariffe il file si ricalcola da solo."""

    def edit(wb):
        det = wb["Dettaglio"]
        # prima riga di dati: Sprint Manager 16 h -> 32 h
        first = next(r for r in range(2, 20) if det[f"C{r}"].value == "Manager")
        det[f"F{first}"] = 32
        wb["Riepilogo"]["B26"] = 0.10  # max sconto
        wb["Riepilogo"]["B35"] = 40000  # prezzo firmato
        wb["Fasi"]["G2"] = 0.10  # contingency sulla fase PM
        tar = wb["Tariffe"]
        manager = next(r for r in range(2, 12) if tar[f"A{r}"].value == "Manager")
        tar[f"B{manager}"] = 2400  # prezzo/giorno raddoppiato

    wb = recalculated(ce_full_xlsx(real, NOW), edit)
    assert formula_errors(wb) == []
    s = wb["Riepilogo"]
    # Manager: 32 h iniziali + 16 h aggiunte = 48 h = 6 giorni a 2.400; gli altri profili invariano
    manager_revenue = 6 * 2400
    others = 47162.50 - 4800 - 0  # ricavi originali senza Manager (4.800)
    pm_phase_before = 10375.0
    pm_phase_after = pm_phase_before - 4800 + manager_revenue
    expected_base = others + manager_revenue
    assert s["B19"].value == pytest.approx(expected_base)
    assert s["B21"].value == pytest.approx(pm_phase_after * 0.10)  # contingency solo sulla fase PM
    assert s["B22"].value == pytest.approx(expected_base + pm_phase_after * 0.10)
    assert s["B27"].value == pytest.approx(s["B25"].value * 0.9)  # sconto del 10%
    assert s["B37"].value == pytest.approx(40000 - s["C22"].value)  # margine firmato
    assert s["C22"].value == pytest.approx(
        20200 + 16 * 100  # +16 ore di Manager a 100 €/ora di costo
    )  # +2 giorni di Manager a 100 €/ora


# ================================================================== scenari del foglio (contingency, sconto, firmato)


@needs_libreoffice
def test_scenario_with_contingency_discount_and_signed_price_matches_the_sheet(api, env):
    sc = FIXTURE["scenarios"]["A"]["sheet"]
    created = create(api, env, code="PS-XLS-A")
    content = real_ce_content(env, 1, max_discount_pct="10", signed_price="40000")
    for phase in content["phases"]:
        phase["contingency_pct"] = "10"
    save(api, env, created["ce"]["id"], content)
    wb = recalculated(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW))
    assert formula_errors(wb) == []
    s = wb["Riepilogo"]
    assert s["B21"].value == pytest.approx(float(sc["E9"]), abs=0.01)  # contingency 4.716,25
    assert s["B25"].value == pytest.approx(
        float(sc["DB_C3"]), abs=0.01
    )  # prezzo progetto 51.878,75
    assert s["B27"].value == pytest.approx(
        float(sc["DB_D3"]), abs=0.001
    )  # prezzo minimo 46.690,875
    assert s["B32"].value == pytest.approx(
        float(sc["M8"]), abs=1e-6
    )  # 64,35 giornate con contingency
    assert s["B33"].value == pytest.approx(float(sc["M9"]), abs=1e-6)  # 12,87 settimane
    assert s["B34"].value == pytest.approx(float(sc["DB_E3"]), abs=1e-4)  # fee media 798,13
    assert s["B36"].value == pytest.approx(float(sc["DB_E4"]), abs=1e-4)  # fee firmata 683,76
    assert s["B37"].value == pytest.approx(float(sc["DB_F12"]), abs=0.01)  # margine firmato 19.800
    assert s["B38"].value == pytest.approx(0.495, abs=1e-9) and s["B39"].value == pytest.approx(
        0.505, abs=1e-9
    )
    assert s["B22"].value == pytest.approx(51878.75, abs=0.01) and s["C22"].value == pytest.approx(
        20200
    )
    assert s["B28"].value == pytest.approx(58.5)  # le giornate "base" non includono la contingency
    assert s["E22"].value == pytest.approx(float(sc["DB_F7"]), abs=1e-6)  # margine % 61,06%


@needs_libreoffice
def test_percent_mode_workbook_matches_the_sheet_scenario(api, env):
    sc = FIXTURE["scenarios"]["B"]["sheet"]
    created = create(api, env, code="PS-XLS-B", planning_mode="percent")
    save(api, env, created["ce"]["id"], percent_ce_content(env, 1))
    wb = recalculated(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW))
    assert formula_errors(wb) == []
    s = wb["Riepilogo"]
    assert s["G22"].value == pytest.approx(float(sc["AB4"]), abs=1e-6)  # 148,85 giornate
    assert s["B22"].value == pytest.approx(float(sc["AB5"]), abs=0.01)  # 125.762,50
    assert s["C22"].value == pytest.approx(float(sc["AB6"]), abs=0.01)  # 56.714,00
    assert s["D22"].value == pytest.approx(float(sc["G3"]), abs=0.01)  # 69.048,50
    st = wb["Staffing mensile"]
    assert [st.cell(row=4, column=6 + i).value for i in range(4)] == [17, 20, 22, 22]
    totals = 6 + 4 + 2  # prima riga del blocco dei totali mensili
    for i, col in enumerate("NOPQ"):
        assert st.cell(row=totals, column=6 + i).value == pytest.approx(
            float(sc[f"{col}4"]), abs=1e-6
        )  # giorni
        assert st.cell(row=totals + 2, column=6 + i).value == pytest.approx(
            float(sc[f"{col}3"]), abs=1e-6
        )  # FTE
        assert st.cell(row=totals + 3, column=6 + i).value == pytest.approx(
            float(sc[f"{col}5"]), abs=0.01
        )  # ricavi
        assert st.cell(row=totals + 4, column=6 + i).value == pytest.approx(
            float(sc[f"{col}6"]), abs=0.01
        )  # costi
    det = wb["Dettaglio"]
    days = {det[f"C{r}"].value: det[f"G{r}"].value for r in range(3, 8) if det[f"C{r}"].value}
    assert days["Manager"] == pytest.approx(float(sc["AB15"])) and days["Senior"] == pytest.approx(
        float(sc["AB16"])
    )
    assert days["Specialist"] == pytest.approx(float(sc["AB17"])) and days[
        "Practice"
    ] == pytest.approx(float(sc["AB18"]))


@needs_libreoffice
def test_percent_workbook_is_live_changing_an_allocation_changes_everything(api, env):
    created = create(api, env, code="PS-XLS-B2", planning_mode="percent")
    save(api, env, created["ce"]["id"], percent_ce_content(env, 1))
    xlsx = ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW)

    def edit(wb):
        st = wb["Staffing mensile"]
        alloc_head = 6 + 4 + 2 + 8  # intestazione del blocco "Allocazione %"
        st[f"F{alloc_head + 1}"] = (
            1.0  # Manager a gennaio: 25% -> 100% (+0,75 x 17 giorni lavorativi)
        )
        st["F3"] = 0  # gennaio senza giorni non lavorativi: 22 giorni lavorativi invece di 17

    wb = recalculated(xlsx, edit)
    assert formula_errors(wb) == []
    assert wb["Staffing mensile"]["F4"].value == 22
    s = wb["Riepilogo"]
    # giorni: gennaio passa da 31,45 a 1,85 FTE -> nuovo totale = 1,85 + 0,75 = 2,60 FTE x 22 giorni
    assert wb["Staffing mensile"].cell(row=12, column=6).value == pytest.approx((1.85 + 0.75) * 22)
    assert s["G22"].value == pytest.approx(148.85 - 31.45 + (1.85 + 0.75) * 22)


# ================================================================== struttura e sicurezza del file (senza LibreOffice)
def test_workbook_structure_types_and_formulas(real):
    wb = load_workbook(BytesIO(ce_full_xlsx(real, NOW)))
    assert wb.sheetnames == [
        "Riepilogo",
        "Dettaglio",
        "Fasi",
        "Profili",
        "Staffing mensile",
        "Tariffe",
        "Note",
    ]
    s = wb["Riepilogo"]
    assert (
        s["B22"].value == "=SUM(B19:B21)" and s["B22"].data_type == "f"
    )  # formule vere, non risultati
    assert s["B27"].value == "=B25*(1-B26)" and s["B34"].value.startswith("=IF(B28=0")
    det = wb["Dettaglio"]
    first = next(r for r in range(2, 20) if det[f"C{r}"].value == "Manager")
    assert (
        det[f"G{first}"].value == f"=F{first}/8" and det[f"J{first}"].value == f"=G{first}*H{first}"
    )
    assert (
        det[f"H{first}"].value.startswith("=VLOOKUP($C") and det[f"F{first}"].value == 16
    )  # ore: dato in blu
    assert isinstance(det[f"F{first}"].value, int | float) and det[
        f"F{first}"
    ].font.color.rgb.endswith("0000FF")
    assert det[f"G{first}"].font.color is None or not str(det[f"G{first}"].font.color.rgb).endswith(
        "0000FF"
    )  # formule nere
    st = wb["Staffing mensile"]
    assert (
        hasattr(st["F1"].value, "year") and st["F1"].number_format == "mmm yyyy"
    )  # mese: vera data
    assert wb["Tariffe"]["B2"].value == 1800 and wb["Tariffe"]["D10"].value == "Sì"  # Esterni
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)
    assert (
        wb.properties.creator == "Portale Conti Economici"
        and wb.properties.title == "CE PS-XLS-REAL v1"
    )


def test_workbook_has_no_cached_errors_or_external_links(real):
    raw = ce_full_xlsx(real, NOW)
    import zipfile

    names = zipfile.ZipFile(BytesIO(raw)).namelist()
    assert not any("externalLink" in n for n in names)
    wb = load_workbook(BytesIO(raw))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value.startswith("#")), (
                    ws.title,
                    cell.coordinate,
                )
                if cell.data_type == "f":
                    assert "[" not in cell.value  # nessun riferimento ad altri file


def test_user_text_that_looks_like_a_formula_stays_text_in_excel(api, env):
    created = create(api, env, code="PS-XLS-EVIL")
    content = simple_content(env, 1, project_name="=1+1")
    content["phases"][0]["lines"][0]["activity"] = '=HYPERLINK("http://evil","clic")'
    content["phases"][0]["name"] = "+cmd|' /C calc'!A0"
    save(api, env, created["ce"]["id"], content)
    wb = load_workbook(BytesIO(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW)))
    det = wb["Dettaglio"]
    for ref, text in (("B3", '=HYPERLINK("http://evil","clic")'), ("A2", "+cmd|' /C calc'!A0")):
        assert det[ref].value == text and det[ref].data_type == "s" and det[ref].quotePrefix is True
    assert wb["Riepilogo"]["B7"].value == "=1+1" and wb["Riepilogo"]["B7"].data_type == "s"


def test_empty_ce_still_produces_a_valid_workbook(api, env):
    created = create(api, env, code="PS-XLS-EMPTY")
    wb = load_workbook(BytesIO(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW)))
    assert wb["Riepilogo"]["B22"].value == "=SUM(B19:B21)" and len(wb.sheetnames) == 7


@needs_libreoffice
def test_empty_ce_workbook_recalculates_without_errors(api, env):
    created = create(api, env, code="PS-XLS-EMPTY2")
    wb = recalculated(ce_full_xlsx(detail_of(api, env, created["ce"]["id"]), NOW))
    assert formula_errors(wb) == []
    s = wb["Riepilogo"]
    assert s["B22"].value == 0 and s["E22"].value == "n/d" and s["B34"].value == "n/d"


def test_reduced_workbook_has_only_values_and_one_visible_sheet(real):
    from app.exports.reduced import reduced_from_detail
    from app.exports.xlsx_export import reduced_xlsx

    wb = load_workbook(BytesIO(reduced_xlsx(reduced_from_detail(real))))
    assert wb.sheetnames == ["Riepilogo"] and wb["Riepilogo"].sheet_state == "visible"
    cells = {
        c.coordinate: c.value
        for row in wb["Riepilogo"].iter_rows()
        for c in row
        if c.value is not None
    }
    assert not any(
        c.data_type == "f" for row in wb["Riepilogo"].iter_rows() for c in row
    )  # nessuna formula
    texts = " ".join(str(v) for v in cells.values())
    assert "PS-XLS-REAL" in texts and "Totale generale" in texts and 47162.5 in cells.values()
    assert 10 in cells.values() and 48.5 in cells.values() and 58.5 in cells.values()
    for forbidden in ("Costi", "Margine", "Tariffe", "Sprint", "Specialist", "Senior", "Manager"):
        assert forbidden not in texts, forbidden
    assert wb.properties.creator == "Portale Conti Economici"
