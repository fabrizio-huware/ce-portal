"""Ricalcolo di un file Excel con LibreOffice: serve a verificare che le formule diano i numeri del portale."""

import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
needs_libreoffice = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice non installato")


def recalculated(xlsx_bytes: bytes, edit=None):
    """Restituisce (cartella di lavoro con i valori ricalcolati da LibreOffice).

    `edit(workbook)` permette di modificare i dati prima del ricalcolo (prova delle formule vive).
    """
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in" / "ce.xlsx"
        out = Path(tmp) / "out"
        src.parent.mkdir()
        out.mkdir()
        if edit:
            wb = load_workbook(BytesIO(xlsx_bytes))
            edit(wb)
            wb.save(src)
        else:
            src.write_bytes(xlsx_bytes)
        subprocess.run(
            [SOFFICE, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--norestore",
             "--convert-to", "xlsx", "--outdir", str(out), str(src)],
            check=True, capture_output=True, timeout=120,
        )  # fmt: skip
        return load_workbook(out / "ce.xlsx", data_only=True)


def formula_errors(wb) -> list[str]:
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    bad.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    return bad
