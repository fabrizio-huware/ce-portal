"""Nessun export "ridotto" deve contenere costi, margini, tariffe, attività, collaboratori o note.

Si guardano i file per intero: celle, nomi dei fogli, proprietà del documento, ogni parte interna dello
zip di Excel, il testo e i metadati del PDF e l'intero CSV.
"""

import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.dash_helpers import build_ce, make_employee

URL = "/api/v1/ce"

# Dati che esistono nel CE ma non devono uscire (cercati come testo, in ogni formato numerico italiano)
SECRET_TEXT = ["SEGRETO-ATTIVITA", "NOTA-INTERNA", "SF-RISERVATA", "BU-RISERVATA", "Riservata", "Senior", "Specialist",
               "Manager", "Collaboratore", "Costo", "Costi", "Margine", "Tariff", "Prezzo/giorno", "C/R", "FTE"]  # fmt: skip
SECRET_NUMBERS = ["8.900,00", "8900,00", "8900.00", "14.100,00", "14100,00", "14100.00", "3.300,00", "3300,00",
                  "5.600,00", "5600,00", "61,30", "0,6130", "0.613", "38,70", "330,00", "280,00", "725,00", "850,00"]  # fmt: skip
ALLOWED = [
    "PS-LEAK-1",
    "Cliente di prova",
    "Progetto riservato ai clienti",
    "Fase Pubblica",
    "23.000,00",
    "23000",
]


@pytest.fixture
def ce(api, env, session):
    ada = make_employee(session, env, "Ada", "Riservata")
    ce_id = build_ce(
        api, env, "PS-LEAK-1", bu="BU-RISERVATA", sf="SF-RISERVATA-123", notes="NOTA-INTERNA-XYZ",
        project="Progetto riservato ai clienti",
        lines=[("Senior", 80, ada, "SEGRETO-ATTIVITA-1"), ("Specialist", 160, None, "SEGRETO-ATTIVITA-2")],
        state="approved",
    )  # fmt: skip
    return ce_id


def xlsx_blobs(raw: bytes) -> list[tuple[str, str]]:
    blobs = []
    with zipfile.ZipFile(BytesIO(raw)) as z:
        for name in z.namelist():
            blobs.append((f"zip:{name}", z.read(name).decode("utf-8", errors="ignore")))
    wb = load_workbook(BytesIO(raw))
    blobs.append(("sheetnames", " ".join(wb.sheetnames)))
    blobs.append(("properties", " ".join(str(v) for v in vars(wb.properties).values() if v)))
    for ws in wb.worksheets:
        assert ws.sheet_state == "visible"
        blobs.append(
            (
                f"cells:{ws.title}",
                " ".join(
                    str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
                ),
            )
        )
    assert not wb.defined_names
    return blobs


def pdf_blobs(raw: bytes) -> list[tuple[str, str]]:
    reader = PdfReader(BytesIO(raw))
    blobs = [("pdf-text", "\n".join(p.extract_text() for p in reader.pages))]
    blobs.append(("pdf-meta", " ".join(str(v) for v in (reader.metadata or {}).values())))
    blobs.append(("pdf-raw", raw.decode("latin-1", errors="ignore")))
    return blobs


def csv_blobs(raw: bytes) -> list[tuple[str, str]]:
    return [("csv", raw.decode("utf-8-sig"))]


def assert_clean(blobs: list[tuple[str, str]], where: str) -> None:
    for label, text in blobs:
        for secret in SECRET_TEXT + SECRET_NUMBERS:
            if label in ("pdf-raw",) and secret in ("Costo", "Costi", "C/R", "FTE"):
                continue  # sequenze brevi che possono comparire per caso nei byte binari del PDF
            assert secret not in text, f"{where}: '{secret}' trovato in {label}"


@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
def test_viewer_summary_export_contains_nothing_sensitive(api, env, ce, fmt):
    raw = api.get(
        f"{URL}/summaries/{ce}/export", headers=env.h(env.viewer), params={"format": fmt}
    ).content
    blobs = {"pdf": pdf_blobs, "xlsx": xlsx_blobs, "csv": csv_blobs}[fmt](raw)
    assert_clean(blobs, f"viewer {fmt}")
    visible = " ".join(
        text for label, text in blobs if label in ("pdf-text", "csv") or label.startswith("cells:")
    )
    for allowed in ("PS-LEAK-1", "Cliente di prova", "Progetto riservato ai clienti", "23"):
        assert allowed in visible, f"{fmt}: manca '{allowed}'"


@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
def test_editor_summary_variant_is_equally_clean(api, env, ce, fmt):
    raw = api.get(
        f"{URL}/{ce}/export",
        headers=env.h(env.presale),
        params={"format": fmt, "variant": "summary"},
    ).content
    blobs = {"pdf": pdf_blobs, "xlsx": xlsx_blobs, "csv": csv_blobs}[fmt](raw)
    assert_clean(blobs, f"editor summary {fmt}")


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_viewer_list_export_contains_nothing_sensitive(api, env, ce, fmt):
    raw = api.get(
        f"{URL}/summaries/export", headers=env.h(env.viewer), params={"format": fmt}
    ).content
    blobs = xlsx_blobs(raw) if fmt == "xlsx" else csv_blobs(raw)
    assert_clean(blobs, f"viewer list {fmt}")
    assert any("PS-LEAK-1" in text for _, text in blobs)


def test_the_check_itself_would_notice_a_leak(api, env, ce):
    """Prova di sensibilità: il PDF completo contiene davvero i dati sensibili, e il controllo li vede."""
    full = api.get(
        f"{URL}/{ce}/export", headers=env.h(env.presale), params={"format": "pdf"}
    ).content
    text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(full)).pages)
    for expected in ("SEGRETO-ATTIVITA-1", "Senior", "8.900,00", "14.100,00"):
        assert expected in text
    with pytest.raises(AssertionError):
        assert_clean(pdf_blobs(full), "completo")
    full_xlsx = api.get(
        f"{URL}/{ce}/export", headers=env.h(env.presale), params={"format": "xlsx"}
    ).content
    with pytest.raises(AssertionError):
        assert_clean(xlsx_blobs(full_xlsx), "completo")
