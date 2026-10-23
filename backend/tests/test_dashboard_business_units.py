"""Dashboard e filtri per business unit: elenco, Portfolio, Carico risorse (sezione e filtro) ed esportazioni."""

from io import BytesIO

from openpyxl import load_workbook

from tests.ce_helpers import save, simple_content
from tests.test_ce_business_units import _employee_in
from tests.test_ce_visibility import PREFIX, URL, Scene, _listed, _switch

WINDOW = "date_from=2026-01-01&date_to=2026-04-30&include_pipeline=true"


def _codes(scene, keys):
    return {scene.code(k) for k in keys}


def _staff(api, env, scene, session):
    """Ada (business unit A) lavora sul CE «a»; Bruno (business unit B) sul CE «b»."""
    ada, bruno = (
        _employee_in(session, env, scene.a, "Ada"),
        _employee_in(session, env, scene.b, "Bruno"),
    )
    session.commit()
    for key, who, person in (("a", env.presale, ada), ("b", env.other, bruno)):
        detail = api.get(f"{URL}/{scene.ce[key]}", headers=env.h(who)).json()
        content = simple_content(env, detail["version"]["revision"])
        content["phases"][0]["lines"][0]["employee_id"] = str(person.id)
        assert save(api, env, scene.ce[key], content, who).status_code == 200
    return ada, bruno


# ------------------------------------------------------------------ elenco
def test_the_list_filters_by_business_unit(api, env, session):
    scene = Scene(api, env, session)
    h = env.h(env.admin)

    def listed(unit):
        resp = api.get(f"{URL}?limit=200&code={PREFIX}&business_unit_id={unit}", headers=h)
        assert resp.status_code == 200
        return {
            i["code"].removeprefix(f"{PREFIX}-").removesuffix("-PJT") for i in resp.json()["items"]
        }

    assert listed(scene.a.id) == _codes(scene, ["a", "ab", "na"])
    assert listed(scene.b.id) == _codes(scene, ["b", "ab", "nb"])
    assert listed(scene.c.id) == set()  # nessun CE coinvolge Digital
    assert api.get(f"{URL}?business_unit_id=non-un-id", headers=h).status_code == 422


def test_the_filter_never_shows_what_the_user_cannot_see(api, env, session):
    scene = Scene(api, env, session)
    _switch(api, env, True)
    resp = api.get(
        f"{URL}?limit=200&code={PREFIX}&business_unit_id={scene.b.id}", headers=env.h(env.presale)
    )
    codes = {
        i["code"].removeprefix(f"{PREFIX}-").removesuffix("-PJT") for i in resp.json()["items"]
    }
    assert codes == {"AB"}  # i CE «b» e «nb» coinvolgono B ma Anna (A) non li vede
    assert _listed(api, env, env.presale) == {"A", "AB", "NA"}  # senza filtro: la regola di sempre


def test_the_list_export_takes_the_same_filter(api, env, session):
    scene = Scene(api, env, session)
    resp = api.get(
        f"{URL}/export?format=csv&code={PREFIX}&business_unit_id={scene.b.id}",
        headers=env.h(env.admin),
    )
    assert resp.status_code == 200
    for key in scene.ce:
        assert (f"{PREFIX}-{scene.code(key)}-PJT" in resp.text) is (key in {"b", "ab", "nb"}), key


# ------------------------------------------------------------------ Portfolio
def test_the_portfolio_filters_by_business_unit(api, env, session):
    scene = Scene(api, env, session)
    resp = api.get(
        f"/api/v1/dashboard/portfolio?business_unit_id={scene.b.id}", headers=env.h(env.admin)
    )
    assert resp.status_code == 200
    for key in scene.ce:
        assert (f"{PREFIX}-{scene.code(key)}-PJT" in resp.text) is (key in {"b", "ab", "nb"}), key
    export = api.get(
        f"/api/v1/dashboard/portfolio/export?format=csv&business_unit_id={scene.b.id}",
        headers=env.h(env.admin),
    )
    assert (
        export.status_code == 200
        and f"{PREFIX}-A-PJT" not in export.text
        and f"{PREFIX}-B-PJT" in export.text
    )


# ------------------------------------------------------------------ Carico risorse
def test_the_resources_have_a_business_unit_section_and_names_on_the_rows(api, env, session):
    scene = Scene(api, env, session)
    _staff(api, env, scene, session)
    body = api.get(f"/api/v1/dashboard/resources?{WINDOW}", headers=env.h(env.admin)).json()
    units = {u["name"]: u for u in body["business_units"]}
    assert set(units) == {scene.a.name, "Cloud"} and all(
        u["kind"] == "business_unit" for u in units.values()
    )
    assert units[scene.a.name]["persons"] == 1 and units["Cloud"]["persons"] == 1
    assert float(units["Cloud"]["total_days"]) > 0
    by_name = {e["name"]: e for e in body["employees"]}
    assert (
        by_name["Ada Test"]["business_unit_name"],
        by_name["Bruno Test"]["business_unit_name"],
    ) == (scene.a.name, "Cloud")
    assert {p["business_unit_name"] for p in body["practices"]} == {scene.a.name, "Cloud"}


def test_the_resources_filter_by_business_unit(api, env, session):
    scene = Scene(api, env, session)
    _staff(api, env, scene, session)
    body = api.get(
        f"/api/v1/dashboard/resources?{WINDOW}&business_unit_id={scene.b.id}",
        headers=env.h(env.admin),
    ).json()
    assert [e["name"] for e in body["employees"]] == ["Bruno Test"]
    assert [u["name"] for u in body["business_units"]] == ["Cloud"]
    assert [p["name"] for p in body["practices"]] == ["Practice di Cloud"]


def test_a_scoped_user_sees_only_the_own_business_unit_section(api, env, session):
    scene = Scene(api, env, session)
    _staff(api, env, scene, session)
    _switch(api, env, True)
    mine = api.get(f"/api/v1/dashboard/resources?{WINDOW}", headers=env.h(env.presale)).json()
    assert [u["name"] for u in mine["business_units"]] == [scene.a.name]
    other = api.get(
        f"/api/v1/dashboard/resources?{WINDOW}&business_unit_id={scene.b.id}",
        headers=env.h(env.presale),
    ).json()
    assert (
        other["employees"] == [] and other["business_units"] == []
    )  # non può guardare un'altra business unit


def test_the_exports_have_the_business_unit_sheet_and_column(api, env, session):
    scene = Scene(api, env, session)
    _staff(api, env, scene, session)
    xlsx = api.get(
        f"/api/v1/dashboard/resources/export?format=xlsx&{WINDOW}", headers=env.h(env.admin)
    )
    wb = load_workbook(BytesIO(xlsx.content))
    assert "Business unit" in wb.sheetnames
    sheet = wb["Business unit"]
    assert sheet["A1"].value == "Business unit" and {sheet["A2"].value, sheet["A3"].value} == {
        scene.a.name,
        "Cloud",
    }
    assert wb["Dati"][1][-1].value == "Business unit"  # in coda
    csv = api.get(
        f"/api/v1/dashboard/resources/export?format=csv&{WINDOW}", headers=env.h(env.admin)
    ).text
    assert csv.splitlines()[0].rstrip(";").endswith("Business unit") and "Cloud" in csv
