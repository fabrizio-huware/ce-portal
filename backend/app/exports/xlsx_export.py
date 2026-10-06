"""Export Excel.

Il file completo del CE contiene **formule vere**: modificando ore, tariffe, contingency, max sconto o
prezzo firmato i totali si ricalcolano. I dati modificabili sono in blu, le formule in nero.
Le viste ridotte (viewer) contengono solo valori.
"""

from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.worksheet import Worksheet

from app.exports.common import (
    FORMULA_TRIGGERS,
    MODE_LABELS,
    STATUS_LABELS,
    fmt_date,
    fmt_datetime,
    yes_no,
)
from app.schemas.ce import CEDetail, CEListItem, ViewerCE, ViewerListItem

FONT = "Arial"
BLUE = "0000FF"
F_NORMAL = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, color=BLUE)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="667085")
FILL_HEAD = PatternFill("solid", fgColor="101828")
FILL_SUB = PatternFill("solid", fgColor="EEF2F6")
FILL_INPUT = PatternFill("solid", fgColor="FFF9DB")
FILL_ALERT = PatternFill("solid", fgColor="FDE2E1")
THIN = Side(style="thin", color="D0D5DD")
BORDER_TOP = Border(top=THIN)

FMT_MONEY = '#,##0.00;(#,##0.00);"-"'
FMT_NUM = '#,##0.00###;(#,##0.00###);"-"'
FMT_PCT = '0.00%;-0.00%;"-"'
FMT_DATE = "dd/mm/yyyy"
FMT_MONTH = "mmm yyyy"
FMT_INT = '#,##0;(#,##0);"-"'


# ====================================================================== celle
def _num(value: Any) -> Any:
    return float(value) if isinstance(value, Decimal) else value


def put(ws: Worksheet, ref: str, value: Any, *, font: Font = F_NORMAL, fmt: str | None = None,
        fill: PatternFill | None = None, align: Alignment | None = None, border: Border | None = None) -> Cell:  # fmt: skip
    """Scrive un valore (mai una formula): un testo che inizia con `=`, `+`, `-`, `@` resta testo."""
    cell = ws[ref]
    cell.value = _num(value)
    if isinstance(value, str):
        cell.data_type = "s"
        if value.startswith(FORMULA_TRIGGERS):
            cell.quotePrefix = True
    _style(cell, font, fmt, fill, align, border)
    return cell


def formula(ws: Worksheet, ref: str, expression: str, *, font: Font = F_NORMAL, fmt: str | None = None,
            fill: PatternFill | None = None, align: Alignment | None = None, border: Border | None = None) -> Cell:  # fmt: skip
    cell = ws[ref]
    cell.value = expression if expression.startswith("=") else "=" + expression
    _style(cell, font, fmt, fill, align, border)
    return cell


def _style(cell: Cell, font: Font, fmt: str | None, fill: PatternFill | None,
           align: Alignment | None, border: Border | None) -> None:  # fmt: skip
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = border


def head(ws: Worksheet, row: int, labels: Sequence[str], first_col: int = 1) -> None:
    for i, label in enumerate(labels):
        put(ws, f"{L(first_col + i)}{row}", label, font=F_HEAD, fill=FILL_HEAD,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))  # fmt: skip


def widths(ws: Worksheet, sizes: dict[str, float]) -> None:
    for letter, size in sizes.items():
        ws.column_dimensions[letter].width = size


def finish(ws: Worksheet, *, landscape: bool = True, freeze: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if freeze:
        ws.freeze_panes = freeze


def _book(title: str) -> Workbook:
    wb = Workbook()
    wb.properties.creator = "Portale Conti Economici"
    wb.properties.title = title
    wb.properties.lastModifiedBy = "Portale Conti Economici"
    return wb


def _bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _frac(percent: Decimal | None) -> float | None:
    """Percentuale del portale (10 = 10%) -> frazione di Excel (0.10)."""
    return None if percent is None else float(Decimal(percent) / 100)


# ====================================================================== CE completo
def ce_full_xlsx(detail: CEDetail, generated_at: datetime) -> bytes:  # noqa: PLR0915
    cal = detail.calculation
    mode = detail.header.planning_mode
    months = [m.month for m in cal.months]
    n_months = len(months)
    rates = detail.rates
    last_rate_row = len(rates) + 1
    lookup = f"Tariffe!$A$2:$D${last_rate_row}"
    calc_by_pos = {(c.phase_index, c.line_index): c for c in cal.lines}
    lines = [(p, l_) for p, phase in enumerate(detail.phases) for l_ in range(len(phase.lines))]
    n_lines = len(lines)

    wb = _book(f"CE {detail.ce.code} v{detail.version.number}")
    ws_sum = wb.active
    ws_sum.title = "Riepilogo"
    ws_det = wb.create_sheet("Dettaglio")
    ws_ph = wb.create_sheet("Fasi")
    ws_pr = wb.create_sheet("Profili")
    ws_st = wb.create_sheet("Staffing mensile")
    ws_ra = wb.create_sheet("Tariffe")
    ws_no = wb.create_sheet("Note")
    ST = "'Staffing mensile'"

    # ---------------------------------------------------------------- posizioni
    det_row: dict[tuple[int, int], int] = {}
    det_head: dict[int, int] = {}
    det_sub: dict[int, int] = {}
    det_range: dict[int, tuple[int, int] | None] = {}
    r = 2
    for p, phase in enumerate(detail.phases):
        det_head[p] = r
        r += 1
        first = r
        for l_ in range(len(phase.lines)):
            det_row[(p, l_)] = r
            r += 1
        det_range[p] = (first, r - 1) if phase.lines else None
        det_sub[p] = r
        r += 1
    det_last = max(r - 1, 2)

    st_first_month_col = 6
    st_total_col = st_first_month_col + n_months
    stf_row = {pos: 7 + i for i, pos in enumerate(lines)}
    st_last_line = 6 + n_lines
    st_tot = st_last_line + 2  # prima riga del blocco dei totali mensili
    alloc_head = st_tot + 8
    alloc_row = {pos: alloc_head + 1 + i for i, pos in enumerate(lines)}

    def mcol(i: int) -> str:
        return L(st_first_month_col + i)

    # ---------------------------------------------------------------- Tariffe
    head(ws_ra, 1, ["Profilo", "Prezzo/giorno (€)", "Costo/giorno (€)", "Esterno"])
    for i, rate in enumerate(rates, start=2):
        put(ws_ra, f"A{i}", rate.profile_name)
        put(ws_ra, f"B{i}", rate.daily_price, font=F_INPUT, fmt=FMT_MONEY)
        put(ws_ra, f"C{i}", rate.daily_cost, font=F_INPUT, fmt=FMT_MONEY)
        put(
            ws_ra,
            f"D{i}",
            yes_no(rate.is_external),
            font=F_INPUT,
            align=Alignment(horizontal="center"),
        )
    put(ws_ra, f"A{last_rate_row + 2}",
        f"Tariffe congelate nel CE (listino {detail.version.rate_year}). Fonte: portale.", font=F_NOTE)  # fmt: skip
    widths(ws_ra, {"A": 28, "B": 18, "C": 18, "D": 10})
    finish(ws_ra, landscape=False, freeze="A2")

    # ---------------------------------------------------------------- Staffing mensile
    put(ws_st, "A1", "Calendario", font=F_BOLD)
    put(ws_st, "A2", "Giorni feriali (lun-ven)")
    put(ws_st, "A3", "Giorni non lavorativi (festività, chiusure)")
    put(ws_st, "A4", "Giorni lavorativi", font=F_BOLD)
    for i, m in enumerate(cal.months):
        c = mcol(i)
        put(ws_st, f"{c}1", m.month, font=F_HEAD, fmt=FMT_MONTH, fill=FILL_HEAD,
            align=Alignment(horizontal="center"))  # fmt: skip
        put(ws_st, f"{c}2", m.weekdays, fmt=FMT_INT)
        put(ws_st, f"{c}3", m.non_working, font=F_INPUT, fmt=FMT_INT, fill=FILL_INPUT)
        formula(ws_st, f"{c}4", f"{c}2-{c}3", font=F_BOLD, fmt=FMT_INT)
    tc = L(st_total_col)
    put(
        ws_st, f"{tc}1", "Totale", font=F_HEAD, fill=FILL_HEAD, align=Alignment(horizontal="center")
    )
    for rr in (2, 3, 4):
        formula(ws_st, f"{tc}{rr}", f"SUM({mcol(0)}{rr}:{mcol(n_months - 1)}{rr})", fmt=FMT_INT,
                font=F_BOLD if rr == 4 else F_NORMAL)  # fmt: skip

    head(ws_st, 6, ["Fase", "Attività", "Profilo", "Prezzo/g (€)", "Costo/g (€)"])
    for i in range(n_months):
        put(ws_st, f"{mcol(i)}6", months[i], font=F_HEAD, fmt=FMT_MONTH, fill=FILL_HEAD,
            align=Alignment(horizontal="center"))  # fmt: skip
    put(
        ws_st,
        f"{tc}6",
        "Giorni totali",
        font=F_HEAD,
        fill=FILL_HEAD,
        align=Alignment(horizontal="center"),
    )
    for (p, l_), sr in stf_row.items():
        phase, line = detail.phases[p], detail.phases[p].lines[l_]
        dr = det_row[(p, l_)]
        put(ws_st, f"A{sr}", phase.name)
        put(ws_st, f"B{sr}", line.activity)
        formula(ws_st, f"C{sr}", f"Dettaglio!C{dr}")
        formula(ws_st, f"D{sr}", f"Dettaglio!H{dr}", fmt=FMT_MONEY)
        formula(ws_st, f"E{sr}", f"Dettaglio!I{dr}", fmt=FMT_MONEY)
        for i, month in enumerate(months):
            ref = f"{mcol(i)}{sr}"
            if mode == "percent":
                formula(ws_st, ref, f"{mcol(i)}{alloc_row[(p, l_)]}*{mcol(i)}$4", fmt=FMT_NUM)
            else:
                days = calc_by_pos[(p, l_)].monthly_days.get(month, Decimal(0))
                put(ws_st, ref, days, font=F_INPUT, fmt=FMT_NUM)
        formula(
            ws_st,
            f"{tc}{sr}",
            f"SUM({mcol(0)}{sr}:{mcol(n_months - 1)}{sr})",
            fmt=FMT_NUM,
            font=F_BOLD,
        )

    totals = [("Giorni", 0), ("Ore", 1), ("FTE (giorni / giorni lavorativi)", 2), ("Ricavi (€)", 3),
              ("Costi (€)", 4), ("Margine (€)", 5)]  # fmt: skip
    put(ws_st, f"A{st_tot - 1}", "Totali mensili", font=F_BOLD)
    for label, k in totals:
        put(ws_st, f"A{st_tot + k}", label, font=F_BOLD)
        for i in range(n_months + 1):
            c = L(st_first_month_col + i)
            rng = f"{c}7:{c}{st_last_line}"
            if k == 0:
                expr = f"SUM({rng})" if n_lines else "0"
            elif k == 1:
                expr = f"{c}{st_tot}*8"
            elif k == 2:
                expr = f'IF({c}4=0,"n/d",{c}{st_tot}/{c}4)'
            elif k == 3:
                expr = f"SUMPRODUCT({rng},$D$7:$D${st_last_line})" if n_lines else "0"
            elif k == 4:
                expr = f"SUMPRODUCT({rng},$E$7:$E${st_last_line})" if n_lines else "0"
            else:
                expr = f"{c}{st_tot + 3}-{c}{st_tot + 4}"
            fmt = {0: FMT_NUM, 1: FMT_NUM, 2: FMT_PCT, 3: FMT_MONEY, 4: FMT_MONEY, 5: FMT_MONEY}[k]
            formula(
                ws_st,
                f"{c}{st_tot + k}",
                expr,
                fmt=fmt,
                font=F_BOLD,
                border=BORDER_TOP if k == 0 else None,
            )
    put(ws_st, f"A{st_tot + 6}",
        "Ricavi e costi mensili escludono la contingency (ricavo aggiuntivo di fase, senza costi).", font=F_NOTE)  # fmt: skip
    if mode == "percent":
        put(ws_st, f"A{alloc_head - 1}", "Allocazione % per mese (modificabile)", font=F_BOLD)
        head(ws_st, alloc_head, ["Fase", "Attività", "Profilo"])
        for i in range(n_months):
            put(ws_st, f"{mcol(i)}{alloc_head}", months[i], font=F_HEAD, fmt=FMT_MONTH, fill=FILL_HEAD,
                align=Alignment(horizontal="center"))  # fmt: skip
        for (p, l_), ar in alloc_row.items():
            phase, line = detail.phases[p], detail.phases[p].lines[l_]
            put(ws_st, f"A{ar}", phase.name)
            put(ws_st, f"B{ar}", line.activity)
            formula(ws_st, f"C{ar}", f"Dettaglio!C{det_row[(p, l_)]}")
            given = {a.month: a.pct for a in line.allocations}
            for i, month in enumerate(months):
                put(ws_st, f"{mcol(i)}{ar}", _frac(given[month]) if month in given else None,
                    font=F_INPUT, fmt=FMT_PCT, fill=FILL_INPUT)  # fmt: skip
        put(ws_st, f"A{alloc_row[lines[-1]] + 2 if lines else alloc_head + 2}",
            "Giorni di una riga in un mese = % x giorni lavorativi del mese (mese intero).", font=F_NOTE)  # fmt: skip
    else:
        put(ws_st, f"A{st_tot + 7}",
            "Modalità ore: la ripartizione delle ore sui mesi (in proporzione ai giorni lavorativi) è "
            "calcolata dal portale e non si aggiorna se modifichi le ore nel foglio Dettaglio.", font=F_NOTE)  # fmt: skip
    widths(ws_st, {"A": 34, "B": 34, "C": 20, "D": 13, "E": 13})
    for i in range(n_months + 1):
        ws_st.column_dimensions[L(st_first_month_col + i)].width = 13
    finish(ws_st, freeze="F7")

    # ---------------------------------------------------------------- Dettaglio
    head(ws_det, 1, ["Fase", "Attività", "Profilo", "Collaboratore", "Project Mgmt", "Ore", "Giorni",
                     "Prezzo/g (€)", "Costo/g (€)", "Ricavo (€)", "Costo (€)", "Margine (€)", "Esterno"])  # fmt: skip
    for p, phase in enumerate(detail.phases):
        hr = det_head[p]
        for c in range(1, 14):
            ws_det.cell(row=hr, column=c).fill = FILL_SUB
        put(ws_det, f"A{hr}", phase.name, font=F_BOLD, fill=FILL_SUB)
        for l_, line in enumerate(phase.lines):
            dr = det_row[(p, l_)]
            put(ws_det, f"A{dr}", phase.name)
            put(ws_det, f"B{dr}", line.activity)
            put(ws_det, f"C{dr}", line.profile_name, font=F_INPUT)
            put(ws_det, f"D{dr}", line.employee_name or "")
            put(ws_det, f"E{dr}", yes_no(line.is_project_management), font=F_INPUT,
                align=Alignment(horizontal="center"))  # fmt: skip
            if mode == "hours":
                put(ws_det, f"F{dr}", line.hours or Decimal(0), font=F_INPUT, fmt=FMT_NUM)
                formula(ws_det, f"G{dr}", f"F{dr}/8", fmt=FMT_NUM)
            else:
                formula(ws_det, f"G{dr}", f"{ST}!{tc}{stf_row[(p, l_)]}", fmt=FMT_NUM)
                formula(ws_det, f"F{dr}", f"G{dr}*8", fmt=FMT_NUM)
            formula(ws_det, f"H{dr}", f"VLOOKUP($C{dr},{lookup},2,FALSE)", fmt=FMT_MONEY)
            formula(ws_det, f"I{dr}", f"VLOOKUP($C{dr},{lookup},3,FALSE)", fmt=FMT_MONEY)
            formula(ws_det, f"J{dr}", f"G{dr}*H{dr}", fmt=FMT_MONEY)
            formula(ws_det, f"K{dr}", f"G{dr}*I{dr}", fmt=FMT_MONEY)
            formula(ws_det, f"L{dr}", f"J{dr}-K{dr}", fmt=FMT_MONEY)
            formula(
                ws_det,
                f"M{dr}",
                f"VLOOKUP($C{dr},{lookup},4,FALSE)",
                align=Alignment(horizontal="center"),
            )
        sr = det_sub[p]
        put(ws_det, f"A{sr}", f"Totale {phase.name}", font=F_BOLD, border=BORDER_TOP)
        for c in "BCDE":
            ws_det[f"{c}{sr}"].border = BORDER_TOP
        rng = det_range[p]
        for c in "FGJKL":
            expr = f"SUM({c}{rng[0]}:{c}{rng[1]})" if rng else "0"
            fmt = FMT_NUM if c in "FG" else FMT_MONEY
            formula(ws_det, f"{c}{sr}", expr, font=F_BOLD, fmt=fmt, border=BORDER_TOP)
        for c in "HIM":
            ws_det[f"{c}{sr}"].border = BORDER_TOP
    widths(ws_det, {"A": 22, "B": 40, "C": 20, "D": 22, "E": 13, "F": 10, "G": 10, "H": 13, "I": 13,
                    "J": 14, "K": 14, "L": 14, "M": 9})  # fmt: skip
    finish(ws_det, freeze="C2")

    # ---------------------------------------------------------------- Fasi
    head(ws_ph, 1, ["Fase", "Ore", "Giorni", "Ricavi (€)", "Costi (€)", "Margine (€)", "Contingency (%)",
                    "Ricavo da contingency (€)", "Ricavi con contingency (€)", "Giorni con contingency"])  # fmt: skip
    n_ph = len(detail.phases)
    for p, phase in enumerate(detail.phases):
        pr = 2 + p
        put(ws_ph, f"A{pr}", phase.name)
        for target, source, fmt in (("B", "F", FMT_NUM), ("C", "G", FMT_NUM), ("D", "J", FMT_MONEY),
                                    ("E", "K", FMT_MONEY), ("F", "L", FMT_MONEY)):  # fmt: skip
            formula(ws_ph, f"{target}{pr}", f"Dettaglio!{source}{det_sub[p]}", fmt=fmt)
        put(
            ws_ph,
            f"G{pr}",
            _frac(phase.contingency_pct),
            font=F_INPUT,
            fmt=FMT_PCT,
            fill=FILL_INPUT,
        )
        formula(ws_ph, f"H{pr}", f"D{pr}*G{pr}", fmt=FMT_MONEY)
        formula(ws_ph, f"I{pr}", f"D{pr}+H{pr}", fmt=FMT_MONEY)
        formula(ws_ph, f"J{pr}", f"C{pr}*(1+G{pr})", fmt=FMT_NUM)
    ph_tot = 2 + n_ph
    put(ws_ph, f"A{ph_tot}", "Totale", font=F_BOLD, border=BORDER_TOP)
    for c, fmt in zip(
        "BCDEFHIJ",
        (FMT_NUM, FMT_NUM, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_NUM),
        strict=True,
    ):
        formula(ws_ph, f"{c}{ph_tot}", f"SUM({c}2:{c}{ph_tot - 1})" if n_ph else "0", font=F_BOLD, fmt=fmt,
                border=BORDER_TOP)  # fmt: skip
    put(ws_ph, f"A{ph_tot + 2}",
        "Contingency = ricavo aggiuntivo della fase (ricavi x %), senza costi aggiuntivi.", font=F_NOTE)  # fmt: skip
    widths(
        ws_ph,
        {"A": 26, "B": 11, "C": 11, "D": 15, "E": 15, "F": 15, "G": 14, "H": 20, "I": 20, "J": 18},
    )
    finish(ws_ph, freeze="B2")

    # ---------------------------------------------------------------- Profili
    head(
        ws_pr,
        1,
        [
            "Profilo",
            "Esterno",
            "Ore",
            "Giorni",
            "Quota ore",
            "Ricavi (€)",
            "Costi (€)",
            "Margine (€)",
        ],
    )
    det_rng = lambda c: f"Dettaglio!${c}$2:${c}${det_last}"  # noqa: E731
    pr_tot = 2 + len(rates)
    for i in range(2, 2 + len(rates)):
        formula(ws_pr, f"A{i}", f"Tariffe!A{i}")
        formula(ws_pr, f"B{i}", f"Tariffe!D{i}", align=Alignment(horizontal="center"))
        formula(ws_pr, f"C{i}", f"SUMIF({det_rng('C')},$A{i},{det_rng('F')})", fmt=FMT_NUM)
        formula(ws_pr, f"D{i}", f"SUMIF({det_rng('C')},$A{i},{det_rng('G')})", fmt=FMT_NUM)
        formula(ws_pr, f"E{i}", f"IF($C${pr_tot}=0,0,C{i}/$C${pr_tot})", fmt=FMT_PCT)
        formula(ws_pr, f"F{i}", f"SUMIF({det_rng('C')},$A{i},{det_rng('J')})", fmt=FMT_MONEY)
        formula(ws_pr, f"G{i}", f"SUMIF({det_rng('C')},$A{i},{det_rng('K')})", fmt=FMT_MONEY)
        formula(ws_pr, f"H{i}", f"F{i}-G{i}", fmt=FMT_MONEY)
    put(ws_pr, f"A{pr_tot}", "Totale", font=F_BOLD, border=BORDER_TOP)
    ws_pr[f"B{pr_tot}"].border = BORDER_TOP
    for c, fmt in zip(
        "CDEFGH", (FMT_NUM, FMT_NUM, FMT_PCT, FMT_MONEY, FMT_MONEY, FMT_MONEY), strict=True
    ):
        formula(
            ws_pr,
            f"{c}{pr_tot}",
            f"SUM({c}2:{c}{pr_tot - 1})",
            font=F_BOLD,
            fmt=fmt,
            border=BORDER_TOP,
        )
    widths(ws_pr, {"A": 28, "B": 10, "C": 11, "D": 11, "E": 11, "F": 15, "G": 15, "H": 15})
    finish(ws_pr, freeze="B2")

    # ---------------------------------------------------------------- Riepilogo
    h, v = detail.header, detail.version
    put(ws_sum, "A1", f"Conto economico {detail.ce.code}", font=F_TITLE)
    put(ws_sum, "A2", f"Generato dal Portale Conti Economici il {fmt_datetime(generated_at)}. "
        "Celle blu su fondo giallo: dati modificabili per simulazioni. Le altre sono formule.", font=F_NOTE)  # fmt: skip
    head(ws_sum, 4, ["Testata", ""])
    approved = (
        f"{fmt_datetime(v.approved_at)} - {v.approved_by.full_name}"
        if v.approved_at and v.approved_by
        else "-"
    )
    info = [("Codice", detail.ce.code), ("Cliente", h.client.name), ("Progetto", h.project_name),
            ("Periodo", f"{fmt_date(h.start_date)} - {fmt_date(h.end_date)}"), ("Modalità", MODE_LABELS[mode]),
            ("Versione", v.number), ("Stato", STATUS_LABELS[v.status]), ("Anno tariffe", v.rate_year),
            ("Approvato", approved), ("Business Unit", h.business_unit or "-"),
            ("Opportunità Salesforce", h.sf_opportunity or "-"), ("Autore", detail.ce.owner.full_name)]  # fmt: skip
    for i, (label, value) in enumerate(info, start=5):
        put(ws_sum, f"A{i}", label, font=F_BOLD)
        put(ws_sum, f"B{i}", value, align=Alignment(horizontal="left"))

    head(
        ws_sum,
        18,
        [
            "Ricavi e costi",
            "Ricavi (€)",
            "Costi (€)",
            "Margine (€)",
            "Margine %",
            "C/R",
            "Giornate",
        ],
    )
    dm, dj, dk, dg = (f"Dettaglio!${c}$2:${c}${det_last}" for c in "MJKG")
    rows = [("Servizi interni", "No"), ("Servizi esterni", "Sì")]
    for i, (label, flag) in enumerate(rows, start=19):
        put(ws_sum, f"A{i}", label)
        formula(ws_sum, f"B{i}", f'SUMIF({dm},"{flag}",{dj})', fmt=FMT_MONEY)
        formula(ws_sum, f"C{i}", f'SUMIF({dm},"{flag}",{dk})', fmt=FMT_MONEY)
        formula(ws_sum, f"G{i}", f'SUMIF({dm},"{flag}",{dg})', fmt=FMT_NUM)
    put(ws_sum, "A21", "Contingency")
    formula(ws_sum, "B21", f"Fasi!H{ph_tot}", fmt=FMT_MONEY)
    put(
        ws_sum, "C21", 0, fmt=FMT_MONEY
    )  # la contingency non ha costi (regola del foglio originale)
    put(ws_sum, "G21", 0, fmt=FMT_NUM)
    put(ws_sum, "A22", "Totale", font=F_BOLD, border=BORDER_TOP)
    formula(ws_sum, "B22", "SUM(B19:B21)", font=F_BOLD, fmt=FMT_MONEY, border=BORDER_TOP)
    formula(ws_sum, "C22", "SUM(C19:C21)", font=F_BOLD, fmt=FMT_MONEY, border=BORDER_TOP)
    formula(ws_sum, "G22", "SUM(G19:G21)", font=F_BOLD, fmt=FMT_NUM, border=BORDER_TOP)
    for i in range(19, 23):
        bold = F_BOLD if i == 22 else F_NORMAL
        border = BORDER_TOP if i == 22 else None
        formula(ws_sum, f"D{i}", f"B{i}-C{i}", font=bold, fmt=FMT_MONEY, border=border)
        formula(ws_sum, f"E{i}", f'IF(B{i}=0,"n/d",D{i}/B{i})', font=bold, fmt=FMT_PCT, border=border,
                align=Alignment(horizontal="right"))  # fmt: skip
        formula(ws_sum, f"F{i}", f'IF(B{i}=0,"n/d",C{i}/B{i})', font=bold, fmt=FMT_PCT, border=border,
                align=Alignment(horizontal="right"))  # fmt: skip

    head(ws_sum, 24, ["Indicatori", "Valore"])
    kpi = [
        (25, "Prezzo progetto (€)", "B22", FMT_MONEY),
        (27, "Prezzo vendita minimo (€)", "B25*(1-B26)", FMT_MONEY),
        (28, "Giornate totali (senza contingency)", "G22", FMT_NUM),
        (
            29,
            "Giornate project management",
            f'SUMIF(Dettaglio!$E$2:$E${det_last},"Sì",{dg})',
            FMT_NUM,
        ),
        (30, "Giornate attività di delivery", "B28-B29", FMT_NUM),
        (31, "Settimane", "B28/5", FMT_NUM),
        (32, "Giornate con contingency", f"Fasi!J{ph_tot}", FMT_NUM),
        (33, "Settimane con contingency", "B32/5", FMT_NUM),
        (34, "Fee media minima (€/giorno)", 'IF(B28=0,"n/d",B27/B28)', FMT_MONEY),
        (36, "Fee media firmata (€/giorno)", 'IF(OR(B35="",B28=0),"n/d",B35/B28)', FMT_MONEY),
        (37, "Margine firmato (€)", 'IF(B35="","n/d",B35-C22)', FMT_MONEY),
        (38, "Margine % firmato", 'IF(OR(B35="",B35=0),"n/d",(B35-C22)/B35)', FMT_PCT),
        (39, "C/R firmato", 'IF(OR(B35="",B35=0),"n/d",C22/B35)', FMT_PCT),
    ]
    for row, label, expr, fmt in kpi:
        put(ws_sum, f"A{row}", label)
        formula(ws_sum, f"B{row}", expr, fmt=fmt, align=Alignment(horizontal="right"))
    put(ws_sum, "A26", "Max sconto (%)")
    put(ws_sum, "B26", _frac(h.max_discount_pct), font=F_INPUT, fmt=FMT_PCT, fill=FILL_INPUT)
    put(ws_sum, "A35", "Prezzo firmato (€)")
    put(ws_sum, "B35", h.signed_price, font=F_INPUT, fmt=FMT_MONEY, fill=FILL_INPUT)
    widths(ws_sum, {"A": 36, "B": 26, "C": 15, "D": 15, "E": 12, "F": 10, "G": 12})
    finish(ws_sum, landscape=False)

    # ---------------------------------------------------------------- Note
    notes = [
        "Come leggere questo file",
        "Celle blu su fondo giallo o con testo blu: dati modificabili (ore, tariffe, contingency, giorni non lavorativi, max sconto, prezzo firmato).",
        "Celle nere: formule. Modificando i dati blu il file si ricalcola; il portale resta la fonte ufficiale del CE.",
        "1 giorno = 8 ore. Ricavo = giorni x prezzo/giorno. Costo = giorni x costo/giorno. Margine = ricavo - costo.",
        "Contingency = ricavo aggiuntivo della fase (ricavi x %), senza costi aggiuntivi e senza giornate nei totali base.",
        "Prezzo vendita minimo = prezzo progetto x (1 - max sconto). Fee media = prezzo / giornate totali (esterni inclusi).",
        f"Tariffe congelate nel CE (listino {v.rate_year}); i giorni lavorativi di ogni mese sono quelli inseriti sul CE (mese intero).",
        "'n/d' indica un valore non calcolabile (divisione per zero o dato mancante).",
        f"Stato del CE al momento dell'esportazione: {STATUS_LABELS[v.status]} (versione {v.number}, revisione {v.revision}).",
    ]
    for i, line in enumerate(notes, start=1):
        put(
            ws_no,
            f"A{i}",
            line,
            font=F_BOLD if i == 1 else F_NORMAL,
            align=Alignment(wrap_text=True, vertical="top"),
        )
    widths(ws_no, {"A": 130})
    finish(ws_no, landscape=False)
    return _bytes(wb)


# ====================================================================== viste ridotte (solo valori)
def reduced_xlsx(ce: ViewerCE) -> bytes:
    wb = _book(f"CE {ce.code} v{ce.version_number} - riepilogo")
    ws = wb.active
    ws.title = "Riepilogo"
    put(ws, "A1", f"Riepilogo {ce.code}", font=F_TITLE)
    info = [("Codice", ce.code), ("Cliente", ce.client_name), ("Progetto", ce.project_name),
            ("Periodo", f"{fmt_date(ce.start_date)} - {fmt_date(ce.end_date)}"), ("Versione", ce.version_number),
            ("Approvato il", fmt_datetime(ce.approved_at) or "-")]  # fmt: skip
    for i, (label, value) in enumerate(info, start=3):
        put(ws, f"A{i}", label, font=F_BOLD)
        put(ws, f"B{i}", value, align=Alignment(horizontal="left"))
    head(ws, 10, ["Giornate", "Valore"])
    for i, (label, value) in enumerate([("Management", ce.days_project_management), ("Delivery", ce.days_delivery),
                                        ("Totali", ce.days_total)], start=11):  # fmt: skip
        put(ws, f"A{i}", label, font=F_BOLD if label == "Totali" else F_NORMAL)
        put(ws, f"B{i}", value, fmt=FMT_NUM)
    head(ws, 15, ["Ricavi per fase", "Ricavi (€)"])
    row = 16
    for phase in ce.phases:
        put(ws, f"A{row}", phase.name)
        put(ws, f"B{row}", phase.revenue, fmt=FMT_MONEY)
        row += 1
    put(ws, f"A{row}", "Totale generale", font=F_BOLD, border=BORDER_TOP)
    put(ws, f"B{row}", ce.total_revenue, font=F_BOLD, fmt=FMT_MONEY, border=BORDER_TOP)
    widths(ws, {"A": 32, "B": 30})
    finish(ws, landscape=False)
    return _bytes(wb)


# ====================================================================== elenchi
def _table(wb: Workbook, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]],
           formats: dict[int, str] | None = None, first: bool = False, size: dict[str, float] | None = None) -> Worksheet:  # fmt: skip
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    head(ws, 1, header)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            put(ws, f"{L(c)}{r}", value, fmt=(formats or {}).get(c))
    ws.auto_filter.ref = f"A1:{L(len(header))}{max(len(rows) + 1, 2)}"
    widths(ws, size or {})
    finish(ws, freeze="A2")
    return ws


def ce_list_xlsx(items: Sequence[CEListItem]) -> bytes:
    wb = _book("Elenco CE")
    header = ["Codice", "Cliente", "Progetto", "Inizio", "Fine", "Modalità", "Versione", "Stato", "Autore",
              "Aggiornato il", "Prezzo (€)", "Margine %", "Giornate"]  # fmt: skip
    rows = [[i.code, i.client.name, i.project_name, i.start_date, i.end_date, MODE_LABELS[i.planning_mode],
             i.version_number, STATUS_LABELS[i.status], i.owner.full_name, fmt_datetime(i.updated_at), i.price,
             i.margin_pct, i.days_total] for i in items]  # fmt: skip
    _table(wb, "Elenco CE", header, rows, {4: FMT_DATE, 5: FMT_DATE, 11: FMT_MONEY, 12: FMT_PCT, 13: FMT_NUM},
           first=True, size={"A": 22, "B": 26, "C": 34, "D": 12, "E": 12, "F": 12, "G": 9, "H": 16, "I": 24, "J": 18,
                             "K": 15, "L": 11, "M": 11})  # fmt: skip
    return _bytes(wb)


def viewer_list_xlsx(items: Sequence[ViewerListItem]) -> bytes:
    wb = _book("Elenco CE approvati")
    header = [
        "Codice",
        "Cliente",
        "Progetto",
        "Inizio",
        "Fine",
        "Versione",
        "Approvato il",
        "Prezzo (€)",
    ]
    rows = [[i.code, i.client_name, i.project_name, i.start_date, i.end_date, i.version_number,
             fmt_datetime(i.approved_at), i.price] for i in items]  # fmt: skip
    _table(wb, "Elenco CE approvati", header, rows, {4: FMT_DATE, 5: FMT_DATE, 8: FMT_MONEY}, first=True,
           size={"A": 22, "B": 26, "C": 34, "D": 12, "E": 12, "F": 9, "G": 18, "H": 15})  # fmt: skip
    return _bytes(wb)
